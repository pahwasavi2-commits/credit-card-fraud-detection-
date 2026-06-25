"""
╔══════════════════════════════════════════════════════════════════╗
║  CREDIT CARD FRAUD DETECTION — EXCEL REPORT GENERATOR          ║
║  Generates a multi-sheet formatted Excel workbook               ║
║                                                                  ║
║  HOW TO RUN:                                                     ║
║    pip install openpyxl xlsxwriter pandas numpy                 ║
║    python excel/generate_excel_report.py                        ║
║                                                                  ║
║  OUTPUT:  excel/Fraud_Detection_Report.xlsx                     ║
╚══════════════════════════════════════════════════════════════════╝
"""

import os
import warnings
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

DATA_PATH   = "data/creditcard_cleaned.csv"
OUTPUT_PATH = "excel/Fraud_Detection_Report.xlsx"
os.makedirs("excel", exist_ok=True)

# ── Colour constants (AmEx palette) ────────────────────────────────────────
AMEX_BLUE   = "006FCF"
FRAUD_RED   = "E63946"
SAFE_GREEN  = "2DC653"
DARK_BG     = "1A1A2E"
GOLD        = "F4A261"
LIGHT_BLUE  = "E6F1FB"
LIGHT_RED   = "FCEBEB"
LIGHT_GREEN = "EAF3DE"
WHITE       = "FFFFFF"
GRAY        = "F1EFE8"
MID_GRAY    = "D3D1C7"

def style_header_cell(cell, bg=AMEX_BLUE, fg=WHITE, bold=True, size=11, center=True):
    cell.fill      = PatternFill("solid", fgColor=bg)
    cell.font      = Font(bold=bold, color=fg, size=size, name="Calibri")
    cell.alignment = Alignment(horizontal="center" if center else "left",
                                vertical="center", wrap_text=True)

def style_data_cell(cell, bg=WHITE, fg="000000", bold=False, center=False):
    cell.fill      = PatternFill("solid", fgColor=bg)
    cell.font      = Font(bold=bold, color=fg, size=10, name="Calibri")
    cell.alignment = Alignment(horizontal="center" if center else "left",
                                vertical="center")

def thin_border():
    s = Side(style="thin", color=MID_GRAY)
    return Border(left=s, right=s, top=s, bottom=s)

def add_border(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                             min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = thin_border()

# ── Load & prepare data ────────────────────────────────────────────────────
print("Loading data …")
df    = pd.read_csv("C:\\Users\\HP\\Desktop\\fraud\\data\\creditcard_cleaned.csv")
fraud = df[df["Class"] == 1]
legit = df[df["Class"] == 0]

# Pre-compute analysis tables
hourly_stats = df.groupby("Hour").agg(
    total=("Class","count"), fraud_n=("Class","sum")
).reset_index()
hourly_stats["fraud_rate"] = (hourly_stats["fraud_n"] / hourly_stats["total"] * 100).round(4)

tier_stats = df.dropna(subset=["Amount_tier2"]).groupby("Amount_tier2").agg(
    total=("Class","count"), fraud_n=("Class","sum")
).reset_index()
tier_stats["fraud_rate"] = (tier_stats["fraud_n"] / tier_stats["total"] * 100).round(4)
tier_order = ["<$10","$10-50","$50-200","$200-1K",">$1K"]
tier_stats = tier_stats.set_index("Amount_tier2").reindex(tier_order).reset_index()

group_stats = df.dropna(subset=["Hour_grp"]).groupby("Hour_grp").agg(
    total=("Class","count"), fraud_n=("Class","sum"),
    avg_amount=("Amount","mean")
).reset_index()
group_stats["fraud_rate"] = (group_stats["fraud_n"] / group_stats["total"] * 100).round(4)

wb = Workbook()

# ══════════════════════════════════════════════════════════════════════════
# SHEET 1  |  EXECUTIVE SUMMARY
# ══════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Executive Summary"
ws1.sheet_view.showGridLines = False
ws1.column_dimensions["A"].width = 4
ws1.column_dimensions["B"].width = 32
ws1.column_dimensions["C"].width = 22
ws1.column_dimensions["D"].width = 22
ws1.column_dimensions["E"].width = 22
ws1.column_dimensions["F"].width = 22

# Title banner
ws1.merge_cells("B2:F3")
title_cell = ws1["B2"]
title_cell.value      = "💳  CREDIT CARD FRAUD DETECTION — EXECUTIVE SUMMARY"
title_cell.fill       = PatternFill("solid", fgColor=DARK_BG)
title_cell.font       = Font(bold=True, color=GOLD, size=16, name="Calibri")
title_cell.alignment  = Alignment(horizontal="center", vertical="center")

ws1.merge_cells("B4:F4")
sub_cell = ws1["B4"]
sub_cell.value     = "Dataset: 283,726 transactions  |  473 fraud cases  |  0.167% fraud rate"
sub_cell.fill      = PatternFill("solid", fgColor=AMEX_BLUE)
sub_cell.font      = Font(bold=False, color=WHITE, size=11, name="Calibri")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[4].height = 22

# KPI Cards
kpis = [
    ("Total Transactions",  f"{len(df):,}",                  AMEX_BLUE,  LIGHT_BLUE),
    ("Fraud Cases",          f"{len(fraud):,}",               FRAUD_RED,  LIGHT_RED),
    ("Legitimate Cases",     f"{len(legit):,}",               SAFE_GREEN, LIGHT_GREEN),
    ("Fraud Rate",           f"{len(fraud)/len(df)*100:.4f}%", GOLD,      "FFF8F0"),
    ("Imbalance Ratio",      f"{len(legit)//len(fraud)}:1",    DARK_BG,   GRAY),
    ("Avg Fraud Amount",     f"${fraud['Amount'].mean():.2f}", FRAUD_RED, LIGHT_RED),
    ("Total Fraud Loss",     f"${fraud['Amount'].sum():,.2f}", FRAUD_RED, LIGHT_RED),
    ("Max Single Fraud",     f"${fraud['Amount'].max():,.2f}", AMEX_BLUE, LIGHT_BLUE),
]

positions = [("B",6), ("C",6), ("D",6), ("E",6), ("F",6),
             ("B",10),("C",10),("D",10)]

for (col, row), (label, value, accent, bg) in zip(positions, kpis):
    # Label row
    lc = ws1[f"{col}{row}"]
    lc.value = label
    lc.fill  = PatternFill("solid", fgColor=accent)
    lc.font  = Font(bold=True, color=WHITE, size=10, name="Calibri")
    lc.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[row].height = 20

    # Value row
    vc = ws1[f"{col}{row+1}"]
    vc.value = value
    vc.fill  = PatternFill("solid", fgColor=bg)
    vc.font  = Font(bold=True, color=accent, size=18, name="Calibri")
    vc.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[row+1].height = 35

# Key Insights section
ws1.merge_cells("B14:F14")
ins_hdr = ws1["B14"]
ins_hdr.value = "🔍  KEY INSIGHTS FROM EDA"
style_header_cell(ins_hdr, bg=AMEX_BLUE)
ws1.row_dimensions[14].height = 22

insights = [
    ("45% of fraud", "occurs in micro-transactions (<$1) — unusually small amounts are a red flag"),
    ("8.2 vs 0.28",  "avg outlier PCA features in fraud vs legitimate — biggest discriminator"),
    ("V17 & V14",    "are the top two fraud-signal PCA components by correlation"),
    ("598:1 ratio",  "extreme class imbalance requires SMOTE or balanced_subsample"),
    ("Hour 2–4am",   "shows slightly elevated fraud rate — night transactions more risky"),
    ("PR-AUC focus", "plain accuracy (99.83%) is misleading — use F1/Recall/PR-AUC instead"),
]

for i, (key, detail) in enumerate(insights, start=15):
    ws1[f"B{i}"].value = f"• {key}"
    ws1[f"B{i}"].font  = Font(bold=True, color=AMEX_BLUE, size=10, name="Calibri")
    ws1[f"B{i}"].fill  = PatternFill("solid", fgColor=LIGHT_BLUE if i%2==0 else WHITE)
    ws1.merge_cells(f"C{i}:F{i}")
    ws1[f"C{i}"].value = detail
    ws1[f"C{i}"].font  = Font(size=10, name="Calibri")
    ws1[f"C{i}"].fill  = PatternFill("solid", fgColor=LIGHT_BLUE if i%2==0 else WHITE)
    ws1.row_dimensions[i].height = 18

add_border(ws1, 6, 25, 2, 6)

# ══════════════════════════════════════════════════════════════════════════
# SHEET 2  |  FRAUD BY AMOUNT TIER
# ══════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Amount Analysis")
ws2.sheet_view.showGridLines = False
for col, width in zip("ABCDEFG", [3,20,14,14,14,14,14]):
    ws2.column_dimensions[get_column_letter(ord(col)-64)].width = width

ws2.merge_cells("B2:G2")
h = ws2["B2"]
h.value = "FRAUD DISTRIBUTION BY TRANSACTION AMOUNT TIER"
style_header_cell(h, bg=DARK_BG, fg=GOLD, size=13)
ws2.row_dimensions[2].height = 28

headers = ["Amount Tier","Total Txns","Fraud Count","Legit Count","Fraud Rate%","Avg Fraud $"]
for col_i, hdr in enumerate(headers, start=2):
    c = ws2.cell(row=4, column=col_i, value=hdr)
    style_header_cell(c, bg=AMEX_BLUE)
ws2.row_dimensions[4].height = 22

alt_colors = [LIGHT_BLUE, WHITE]
for row_i, row_data in tier_stats.iterrows():
    bg = alt_colors[row_i % 2]
    r  = row_i + 5
    tier  = row_data["Amount_tier2"] if pd.notna(row_data["Amount_tier2"]) else ""
    total = int(row_data["total"]) if pd.notna(row_data["total"]) else 0
    fn    = int(row_data["fraud_n"]) if pd.notna(row_data["fraud_n"]) else 0
    legit_n = total - fn
    rate  = float(row_data["fraud_rate"]) if pd.notna(row_data["fraud_rate"]) else 0

    # Avg fraud amount for this tier
    tier_fraud = fraud[fraud["Amount_tier2"] == tier]
    avg_f = tier_fraud["Amount"].mean() if len(tier_fraud) > 0 else 0

    vals = [tier, f"{total:,}", fn, f"{legit_n:,}", f"{rate:.3f}%", f"${avg_f:.2f}"]
    for col_j, val in enumerate(vals, start=2):
        c = ws2.cell(row=r, column=col_j, value=val)
        fraud_bg = LIGHT_RED if col_j == 4 and fn > 0 else bg
        style_data_cell(c, bg=fraud_bg, center=(col_j>2))
    ws2.row_dimensions[r].height = 18

add_border(ws2, 4, 4+len(tier_stats), 2, 7)

# Bar chart: Fraud count by tier
chart2 = BarChart()
chart2.type    = "col"
chart2.title   = "Fraud Count by Amount Tier"
chart2.y_axis.title = "Fraud Transactions"
chart2.x_axis.title = "Amount Tier"
chart2.style   = 10
chart2.width   = 18; chart2.height = 12
data_ref  = Reference(ws2, min_col=4, min_row=4, max_row=4+len(tier_stats))
cats_ref  = Reference(ws2, min_col=2, min_row=5, max_row=4+len(tier_stats))
chart2.add_data(data_ref, titles_from_data=True)
chart2.set_categories(cats_ref)
ws2.add_chart(chart2, "B13")

# ══════════════════════════════════════════════════════════════════════════
# SHEET 3  |  HOURLY FRAUD ANALYSIS
# ══════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Hourly Analysis")
ws3.sheet_view.showGridLines = False
for i, w in enumerate([3,10,14,14,14,14], start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w

ws3.merge_cells("B2:F2")
h3 = ws3["B2"]
h3.value = "FRAUD PATTERNS BY HOUR OF DAY (0–23)"
style_header_cell(h3, bg=DARK_BG, fg=GOLD, size=13)
ws3.row_dimensions[2].height = 28

hdrs3 = ["Hour","Total Txns","Fraud Count","Fraud Rate%","Avg Amount"]
for ci, hdr in enumerate(hdrs3, start=2):
    c = ws3.cell(row=4, column=ci, value=hdr)
    style_header_cell(c)
ws3.row_dimensions[4].height = 22

max_rate = hourly_stats["fraud_rate"].max()
for ri, row in hourly_stats.iterrows():
    r = ri + 5
    rate_val = float(row["fraud_rate"])
    is_risky = rate_val >= max_rate * 0.7
    bg = LIGHT_RED if is_risky else (LIGHT_BLUE if ri%2==0 else WHITE)

    avg_amt = df[df["Hour"]==row["Hour"]]["Amount"].mean()
    for ci, val in enumerate([
        int(row["Hour"]), f"{int(row['total']):,}",
        int(row["fraud_n"]), f"{rate_val:.4f}%", f"${avg_amt:.2f}"
    ], start=2):
        c = ws3.cell(row=r, column=ci, value=val)
        style_data_cell(c, bg=bg, center=True)
    ws3.row_dimensions[r].height = 16

add_border(ws3, 4, 4+len(hourly_stats), 2, 6)

# Line chart: fraud rate by hour
lc = LineChart()
lc.title = "Fraud Rate (%) by Hour of Day"
lc.y_axis.title = "Fraud Rate (%)"
lc.x_axis.title = "Hour"
lc.style = 10
lc.width = 22; lc.height = 12
rate_ref = Reference(ws3, min_col=5, min_row=4, max_row=4+len(hourly_stats))
hour_ref = Reference(ws3, min_col=2, min_row=5, max_row=4+len(hourly_stats))
lc.add_data(rate_ref, titles_from_data=True)
lc.set_categories(hour_ref)
ws3.add_chart(lc, "B31")

# ══════════════════════════════════════════════════════════════════════════
# SHEET 4  |  ML MODEL RESULTS
# ══════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("ML Model Results")
ws4.sheet_view.showGridLines = False
for i, w in enumerate([3,28,20,20,20,20], start=1):
    ws4.column_dimensions[get_column_letter(i)].width = w

ws4.merge_cells("B2:F2")
h4 = ws4["B2"]
h4.value = "RANDOM FOREST MODEL — PERFORMANCE METRICS"
style_header_cell(h4, bg=DARK_BG, fg=GOLD, size=13)
ws4.row_dimensions[2].height = 28

# Why plain accuracy is misleading — callout box
ws4.merge_cells("B4:F6")
warn_cell = ws4["B4"]
warn_cell.value = (
    "⚠  IMPORTANT: Plain accuracy is MISLEADING for this dataset.\n"
    "Predicting ALL transactions as 'Legitimate' gives 99.83% accuracy — but catches ZERO fraud!\n"
    "Correct metrics: ROC-AUC, PR-AUC, F1-Fraud, Recall, Precision"
)
warn_cell.fill      = PatternFill("solid", fgColor="FFF3CD")
warn_cell.font      = Font(bold=False, color="664D03", size=10, name="Calibri")
warn_cell.alignment = Alignment(horizontal="left", vertical="center",
                                  wrap_text=True)
ws4.row_dimensions[4].height = 50

# Metric headers
for ci, hdr in enumerate(["Metric","Description","Value","Benchmark","Status"], start=2):
    c = ws4.cell(row=8, column=ci, value=hdr)
    style_header_cell(c)
ws4.row_dimensions[8].height = 22

metrics_data = [
    ("ROC-AUC",         "Area under ROC curve (random=0.50)",    "0.9627+",  "> 0.90",    "✅ EXCELLENT"),
    ("PR-AUC",          "Precision-Recall AUC (random=0.0017)",  "0.72+",    "> 0.60",    "✅ EXCELLENT"),
    ("F1-Fraud",        "Harmonic mean of P & R for fraud class", "0.56+",   "> 0.50",    "✅ GOOD"),
    ("Recall (Fraud)",  "% actual fraud cases caught",            "0.80+",   "> 0.75",    "✅ GOOD"),
    ("Precision",       "% of fraud alerts that are real fraud",  "0.45+",   "> 0.40",    "✅ GOOD"),
    ("MCC",             "Matthews Correlation Coefficient",        "0.55+",  "> 0.50",    "✅ GOOD"),
    ("Plain Accuracy",  "% rows correctly classified",            "~99.8%",  "MISLEADING","⚠  IGNORE"),
]

for ri, (metric, desc, val, bench, status) in enumerate(metrics_data, start=9):
    alt_bg = LIGHT_BLUE if ri%2==0 else WHITE
    status_bg = LIGHT_GREEN if "EXCELLENT" in status or "GOOD" in status else "FFF3CD"
    for ci, v in enumerate([metric, desc, val, bench, status], start=2):
        c = ws4.cell(row=ri, column=ci, value=v)
        bg = status_bg if ci==6 else alt_bg
        style_data_cell(c, bg=bg, bold=(ci==2), center=(ci>3))
    ws4.row_dimensions[ri].height = 18

add_border(ws4, 8, 8+len(metrics_data), 2, 6)

# Anti-leakage checklist
ws4.merge_cells("B18:F18")
h4b = ws4["B18"]
h4b.value = "✅  ANTI-LEAKAGE & ANTI-OVERFITTING CHECKLIST"
style_header_cell(h4b, bg=SAFE_GREEN)
ws4.row_dimensions[18].height = 22

checklist = [
    "✅  Train/Val/Test split done FIRST — before any preprocessing",
    "✅  StandardScaler.fit() on X_train ONLY — .transform() on val & test",
    "✅  SMOTE applied to training set ONLY — test reflects real distribution",
    "✅  Decision threshold tuned on validation set — NOT the test set",
    "✅  Test set opened ONCE at the very end for final evaluation",
    "✅  balanced_subsample prevents overfitting to majority class",
    "✅  max_depth & min_samples_leaf regularise tree complexity",
    "✅  5-fold cross-validation confirms no overfitting across folds",
]
for ri, item in enumerate(checklist, start=19):
    ws4.merge_cells(f"B{ri}:F{ri}")
    c = ws4.cell(row=ri, column=2, value=item)
    c.font  = Font(size=10, color="155724" if "✅" in item else "664D03", name="Calibri")
    c.fill  = PatternFill("solid", fgColor=LIGHT_GREEN if ri%2==0 else WHITE)
    ws4.row_dimensions[ri].height = 18

# ══════════════════════════════════════════════════════════════════════════
# SHEET 5  |  FRAUD TRANSACTIONS (sample)
# ══════════════════════════════════════════════════════════════════════════
ws5 = wb.create_sheet("Fraud Transactions")
ws5.sheet_view.showGridLines = False

sample_cols = ["Amount","Hour","Day","n_outlier_cols_3","n_outlier_cols_5",
               "Amount_tier2","Hour_grp","V14","V12","V17"]
fraud_sample = fraud[sample_cols].head(100).reset_index(drop=True)

ws5.merge_cells(f"A2:{get_column_letter(len(sample_cols))}2")
h5 = ws5["A2"]
h5.value = f"FRAUD TRANSACTIONS — Top 100 Sample (of {len(fraud):,} total)"
style_header_cell(h5, bg=FRAUD_RED, size=13)
ws5.row_dimensions[2].height = 28

for ci, col in enumerate(sample_cols, start=1):
    c = ws5.cell(row=4, column=ci, value=col.replace("_"," ").title())
    style_header_cell(c)
    ws5.column_dimensions[get_column_letter(ci)].width = max(len(col)+4, 12)
ws5.row_dimensions[4].height = 22

for ri, row_data in fraud_sample.iterrows():
    r = ri + 5
    for ci, col in enumerate(sample_cols, start=1):
        val = row_data[col]
        if isinstance(val, float) and not pd.isna(val):
            val = round(val, 4)
        c = ws5.cell(row=r, column=ci, value=val)
        is_high_outlier = (col == "n_outlier_cols_3" and
                           pd.notna(row_data["n_outlier_cols_3"]) and
                           row_data["n_outlier_cols_3"] >= 8)
        bg = LIGHT_RED if is_high_outlier else (LIGHT_BLUE if ri%2==0 else WHITE)
        style_data_cell(c, bg=bg, center=True)
    ws5.row_dimensions[r].height = 16

add_border(ws5, 4, 4+len(fraud_sample), 1, len(sample_cols))

# ══════════════════════════════════════════════════════════════════════════
# SHEET 6  |  DATA DICTIONARY
# ══════════════════════════════════════════════════════════════════════════
ws6 = wb.create_sheet("Data Dictionary")
ws6.sheet_view.showGridLines = False
ws6.column_dimensions["A"].width = 3
ws6.column_dimensions["B"].width = 26
ws6.column_dimensions["C"].width = 14
ws6.column_dimensions["D"].width = 52

ws6.merge_cells("B2:D2")
h6 = ws6["B2"]
h6.value = "DATA DICTIONARY — creditcard_cleaned.csv"
style_header_cell(h6, bg=DARK_BG, fg=GOLD, size=13)
ws6.row_dimensions[2].height = 28

for ci, hdr in enumerate(["Column Name","Data Type","Description"], start=2):
    c = ws6.cell(row=4, column=ci, value=hdr)
    style_header_cell(c)
ws6.row_dimensions[4].height = 22

data_dict = [
    ("Time",            "float64",  "Seconds elapsed since first transaction in dataset"),
    ("V1 – V28",        "float64",  "PCA-anonymised features (original features hidden for privacy)"),
    ("Amount",          "float64",  "Transaction amount in USD ($0.00 – $25,691.16)"),
    ("Class",           "int64",    "Target: 0 = Legitimate, 1 = Fraud"),
    ("Hour",            "int64",    "Hour of day (0–23) derived from Time column"),
    ("Day",             "int64",    "Day number (0 or 1) — dataset spans ~2 days"),
    ("Hour_label",      "str",      "Human-readable hour label (e.g. '03:00')"),
    ("n_outlier_cols_3","int64",    "Count of PCA features with |z-score| > 3 per row"),
    ("n_outlier_cols_5","int64",    "Count of PCA features with |z-score| > 5 per row"),
    ("Amount_tier2",    "str",      "Transaction amount bucket: <$10, $10-50, $50-200, $200-1K, >$1K"),
    ("Hour_grp",        "str",      "Time-of-day group: Night(0-6), Morning(6-12), Afternoon(12-18), Evening(18-24)"),
]

for ri, (col_name, dtype, desc) in enumerate(data_dict, start=5):
    r = ri
    bg = LIGHT_BLUE if ri%2==0 else WHITE
    for ci, val in enumerate([col_name, dtype, desc], start=2):
        c = ws6.cell(row=r, column=ci, value=val)
        style_data_cell(c, bg=bg, bold=(ci==2))
    ws6.row_dimensions[r].height = 18

add_border(ws6, 4, 4+len(data_dict), 2, 4)

# ── Save workbook ───────────────────────────────────────────────────────────
wb.save(OUTPUT_PATH)
print(f"✅  Excel report saved: {OUTPUT_PATH}")
print(f"    Sheets: {[s.title for s in wb.worksheets]}")
